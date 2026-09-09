"""Caso 2: consulta manual en ONPE, registro y descarga de Excel."""

from io import BytesIO
import os
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from records import FIELDS, HEADERS, ValidationError, add_record, initialize, list_records, remove_record

ONPE_URL = "https://consultaelectoral.onpe.gob.pe/inicio"


def create_app(config=None):
    app = Flask(__name__)
    app.config.update(
        MAX_CONTENT_LENGTH=16384,
        DATABASE=str(Path(os.environ.get("DATA_DIR", str(Path(__file__).parent / "data"))) / "registros.sqlite3"),
    )
    if config:
        app.config.update(config)
    initialize(app.config["DATABASE"])

    @app.after_request
    def disable_cache(response):
        response.headers["Cache-Control"] = "no-store"
        return response

    @app.get("/")
    def index():
        return render_template("index.html", onpe_url=ONPE_URL)

    @app.get("/health")
    def health():
        list_records(app.config["DATABASE"])
        return jsonify(status="ok")

    @app.get("/api/registros")
    def get_records():
        return jsonify(records=list_records(app.config["DATABASE"]))

    @app.post("/api/registros")
    def post_record():
        try:
            record = add_record(app.config["DATABASE"], request.get_json(silent=True))
        except ValidationError as error:
            return jsonify(error=str(error)), 400
        return jsonify(record=record), 201

    @app.delete("/api/registros/<dni>")
    def delete_record(dni):
        if not remove_record(app.config["DATABASE"], dni):
            return jsonify(error="El registro ya no existe."), 404
        return jsonify(deleted=True)

    @app.get("/exportar.xlsx")
    def export_excel():
        rows = list_records(app.config["DATABASE"])
        if not rows:
            return jsonify(error="Agrega al menos un registro antes de exportar."), 400
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Consulta electoral"
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17645D")
        for record in rows:
            sheet.append([record[field] for field in FIELDS])
            for cell in sheet[sheet.max_row]:
                # Conserva DNI con ceros iniciales y evita interpretar entradas como fórmulas.
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, width in enumerate((14, 20, 36, 22, 22, 22, 50), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="consulta-electoral.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.errorhandler(413)
    def too_large(_error):
        return jsonify(error="La solicitud supera el tamaño permitido."), 413

    return app


if __name__ == "__main__":
    create_app().run(host="0.0.0.0", port=5000, debug=False, threaded=True)
